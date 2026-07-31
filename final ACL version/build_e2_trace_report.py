from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = RGBColor(24, 37, 52)
BLUE = RGBColor(32, 78, 121)
DARK_BLUE = RGBColor(23, 54, 83)
MUTED = RGBColor(91, 103, 116)
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
PALE_GOLD = "FFF5D9"
WHITE = RGBColor(255, 255, 255)

STAGE_LABELS = {
    "audience_expectation": "Audience expectation",
    "opposition_candidates": "Opposition candidates",
    "theory_gate": "Theory gate",
    "candidate_selection": "Candidate selection",
    "gtvh_plan": "GTVH plan",
    "variants": "Controlled joke variants",
    "blind_reconstruction": "Blind reconstruction",
    "pairwise_selection": "Pairwise selection",
}

STAGE_PURPOSES = {
    "audience_expectation": "Recovers the audience’s default Script A and its expected propositions.",
    "opposition_candidates": "Constructs eight explicit Script B candidates with anchors and logical mechanisms.",
    "theory_gate": "Applies the non-compensatory SSTH/GTVH validity gate to every candidate.",
    "candidate_selection": "Selects one theory-valid opposition by immutable candidate ID.",
    "gtvh_plan": "Expands the selected opposition across the six GTVH knowledge resources.",
    "variants": "Realizes the plan as three setup/punchline-separated Dutch jokes.",
    "blind_reconstruction": "Checks whether the finished jokes independently reveal both scripts and their resolution.",
    "pairwise_selection": "Compares every eligible pair and selects the strongest final joke.",
}


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run_font(run, name: str, size: float, color=INK, bold=False, italic=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade_paragraph(paragraph, fill: str, border_color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)

    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "14")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border_color)
    borders.append(left)
    p_pr.append(borders)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Title": (30, DARK_BLUE, 0, 8),
        "Subtitle": (14, MUTED, 0, 18),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("ACL FINAL PIPELINE  /  E2 TRACE REPORT")
    set_run_font(run, "Calibri", 8.5, MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("PAGE ")
    set_run_font(run, "Calibri", 8.5, MUTED)
    add_field(p, "PAGE")
    run = p.add_run("  ·  GPT-5.4 MINI")
    set_run_font(run, "Calibri", 8.5, MUTED)


def add_cover(doc: Document, result: dict) -> None:
    for _ in range(3):
        doc.add_paragraph()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("EXPERIMENT TRACE")
    set_run_font(run, "Calibri", 10, BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("E2 Validated GTVH Pipeline")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Complete prompts, intermediary outputs, and final selection")

    metadata = [
        ("Model", result["usage"]["model"]),
        ("Topic", result["request"]["topic"]),
        ("Audience", result["request"]["audience"]),
        ("Format", result["request"]["joke_format"]),
        ("Token usage", f'{result["usage"]["input_tokens"]:,} input  ·  '
                        f'{result["usage"]["output_tokens"]:,} output  ·  '
                        f'{result["usage"]["total_tokens"]:,} total'),
    ]
    table = doc.add_table(rows=0, cols=2)
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], PALE_BLUE)
        for cell in cells:
            cell.text = ""
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(label.upper()), "Calibri", 9, BLUE, bold=True)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), "Calibri", 10.5, INK)
    set_fixed_table_geometry(table, [1900, 7460])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("SELECTED JOKE")
    set_run_font(run, "Calibri", 9, BLUE, bold=True)

    callout = doc.add_paragraph()
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(0)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.2
    shade_paragraph(callout, PALE_GOLD, "D29B20")
    run = callout.add_run(result["joke"])
    set_run_font(run, "Calibri", 13, DARK_BLUE, bold=True)

    trace = result["metadata"]["gtvh_trace"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(
        f'Selected candidate {trace["selected_candidate_id"]}  →  '
        f'selected variant {trace["selected_variant_id"]}  →  '
        f'{len(trace["eligible_variant_ids"])} eligible variants'
    )
    set_run_font(run, "Calibri", 9.5, MUTED, italic=True)
    doc.add_page_break()


def add_overview(doc: Document, result: dict, stages: list[str]) -> None:
    doc.add_heading("Run overview", level=1)
    p = doc.add_paragraph(
        "This report preserves the exact model-facing prompt and exact structured response "
        "for each completed stage. The ordering below follows the runtime trace."
    )
    p.paragraph_format.space_after = Pt(12)

    doc.add_heading("Pipeline path", level=2)
    for index, stage in enumerate(stages, start=1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        lead = p.add_run(f"{STAGE_LABELS[stage]}. ")
        lead.bold = True
        p.add_run(STAGE_PURPOSES[stage])

    trace = result["metadata"]["gtvh_trace"]
    doc.add_heading("Outcome at a glance", level=2)
    outcome = [
        ("Theory gate", f'{len(trace["approved_candidate_ids"])} of 8 candidates approved'),
        ("Repair stage", "Not triggered"),
        ("Candidate selected", trace["selected_candidate_id"]),
        ("Blind reconstruction", f'{len(trace["eligible_variant_ids"])} of 3 variants eligible'),
        ("Final selection", trace["selected_variant_id"]),
        ("Warnings", "None" if not result["warnings"] else "; ".join(result["warnings"])),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_repeat_table_header(table.rows[0])
    table.rows[0].cells[0].text = "Checkpoint"
    table.rows[0].cells[1].text = "Result"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, PALE_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, "Calibri", 9.5, BLUE, bold=True)
    for label, value in outcome:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, "Calibri", 9.5, INK)
    set_fixed_table_geometry(table, [2900, 6460])

    doc.add_heading("Reading convention", level=2)
    p = doc.add_paragraph(
        "Prompt blocks use a blue rule and cool-gray background. Response blocks use a "
        "darker rule and warmer-gray background. JSON is pretty-printed for legibility; "
        "field names and values are unchanged."
    )
    p.paragraph_format.space_after = Pt(0)


def add_trace_block(doc: Document, text: str, *, kind: str) -> None:
    fill = PALE_GRAY if kind == "prompt" else "F3F1EC"
    border = "5B8DB8" if kind == "prompt" else "6D6258"
    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.05)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = False
        shade_paragraph(p, fill, border)
        run = p.add_run(line if line else " ")
        set_run_font(run, "Aptos Mono", 8.0, INK)
        if index == 0:
            p.paragraph_format.space_before = Pt(3)
        if index == len(lines) - 1:
            p.paragraph_format.space_after = Pt(4)


def split_prompts(combined: str) -> dict[str, str]:
    parts = re.split(r"^## ([a-z0-9_]+)\n", combined, flags=re.MULTILINE)
    prompts: dict[str, str] = {}
    for i in range(1, len(parts), 2):
        prompt = re.sub(r"\n+---\s*$", "", parts[i + 1].rstrip())
        prompts[parts[i]] = prompt
    return prompts


def prettify_response(raw: str) -> str:
    try:
        return json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return raw


def add_stage(
    doc: Document,
    stage: str,
    index: int,
    total: int,
    prompt: str,
    response: str,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"STAGE {index:02d} / {total:02d}")
    set_run_font(run, "Calibri", 9, MUTED, bold=True)

    doc.add_heading(STAGE_LABELS[stage], level=1)
    p = doc.add_paragraph(STAGE_PURPOSES[stage])
    p.paragraph_format.space_after = Pt(12)

    doc.add_heading("Prompt sent to gpt-5.4-mini", level=2)
    add_trace_block(doc, prompt, kind="prompt")

    response_heading = doc.add_heading("Structured response", level=2)
    response_heading.paragraph_format.page_break_before = True
    add_trace_block(doc, prettify_response(response), kind="response")


def build_report(input_path: Path, output_path: Path) -> None:
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    if not isinstance(payload, list) or len(payload) != 1:
        raise ValueError("Expected one pipeline result in the input trace.")
    result = payload[0]
    stages = result["metadata"]["stages"]
    prompts = split_prompts(result["prompt"])
    responses = json.loads(result["raw_response"])

    missing_prompts = [stage for stage in stages if stage not in prompts]
    missing_responses = [stage for stage in stages if stage not in responses]
    if missing_prompts or missing_responses:
        raise ValueError(
            f"Trace mismatch: missing prompts={missing_prompts}, missing responses={missing_responses}"
        )

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc, result)
    add_overview(doc, result, stages)
    for index, stage in enumerate(stages, start=1):
        add_stage(doc, stage, index, len(stages), prompts[stage], responses[stage])

    end_heading = doc.add_heading("End of trace", level=1)
    end_heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph(
        "All eight completed runtime stages are included. No opposition-repair stage "
        "appears because every initial candidate passed the theory gate."
    )
    p.paragraph_format.space_after = Pt(12)
    source = doc.add_paragraph()
    source.add_run("Source trace: ").bold = True
    source.add_run(input_path.name)
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_e2_trace_report.py INPUT_JSON OUTPUT_DOCX")
    build_report(Path(sys.argv[1]), Path(sys.argv[2]))
